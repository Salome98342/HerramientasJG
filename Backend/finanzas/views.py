from datetime import datetime, time, timedelta
from decimal import Decimal
from io import BytesIO

from django.db.models import ExpressionWrapper, F, Sum
from django.db.models.functions import TruncDate
from django.http import HttpResponse
from django.utils import timezone
from django.utils.dateparse import parse_date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from rest_framework import generics, status
from rest_framework.exceptions import ValidationError
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from cajas.models import TurnoCaja
from cajas.services import resumen_turno
from core.permissions import IsAdmin, IsAdminOrCajero
from inventario.models import ProductoVenta
from ventas.models import Venta

from .models import AlertaLeida, CategoriaGasto, Gasto
from .reports import (
    alquileres_vencidos,
    articulos_mas_alquilados,
    cartera_credito_por_cliente,
    estado_caja_por_turno,
    gastos_por_categoria,
    ingresos_egresos_por_medio,
    productos_bajo_minimo,
    productos_mas_vendidos,
    resumen_periodo,
    serie_diaria,
)
from .serializers import (
    CategoriaGastoSerializer,
    GastoSerializer,
    RegistrarGastoSerializer,
)
from .services import registrar_gasto


def _rango_fechas(request, *, requerido=False):
    desde_raw = request.query_params.get('desde')
    hasta_raw = request.query_params.get('hasta')
    if requerido and (not desde_raw or not hasta_raw):
        raise ValidationError({'detail': 'Los parámetros desde y hasta son obligatorios.'})
    desde_dia = parse_date(desde_raw) if desde_raw else timezone.localdate()
    hasta_dia = parse_date(hasta_raw) if hasta_raw else timezone.localdate()
    if desde_dia is None or hasta_dia is None:
        raise ValidationError({'detail': 'Usa fechas válidas en formato AAAA-MM-DD.'})
    if desde_dia > hasta_dia:
        raise ValidationError({'detail': 'La fecha desde no puede ser posterior a hasta.'})
    tz = timezone.get_current_timezone()
    desde = timezone.make_aware(datetime.combine(desde_dia, time.min), tz)
    hasta = timezone.make_aware(
        datetime.combine(hasta_dia + timedelta(days=1), time.min), tz,
    )
    return desde_dia, hasta_dia, desde, hasta


def _monto(valor):
    if valor is None:
        return None
    return f'{valor:.2f}' if isinstance(valor, Decimal) else str(valor)


def _reporte(negocio, desde, hasta):
    resumen = resumen_periodo(negocio, desde, hasta)
    return {
        'resumen': {key: _monto(value) for key, value in resumen.items()},
        'serie_diaria': [
            {
                'dia': row['dia'].isoformat(),
                'ventas': _monto(row['ventas']),
                'gastos': _monto(row['gastos']),
            }
            for row in serie_diaria(negocio, desde, hasta)
        ],
        'gastos_por_categoria': [
            {
                'categoria_id': row['categoria_id'],
                'categoria': row['categoria__nombre'],
                'total': _monto(row['total']),
            }
            for row in gastos_por_categoria(negocio, desde, hasta)
        ],
        'ingresos_egresos_por_medio': [
            {
                'medio_pago': row['medio_pago'],
                'ingresos': _monto(row['ingresos']),
                'egresos': _monto(row['egresos']),
            }
            for row in ingresos_egresos_por_medio(negocio, desde, hasta)
        ],
        'productos_mas_vendidos': [
            {
                'producto_id': row['producto_id'],
                'referencia': row['producto__referencia'],
                'producto': row['producto__nombre'],
                'cantidad': _monto(row['cantidad']),
                'ingresos': _monto(row['ingresos']),
            }
            for row in productos_mas_vendidos(negocio, desde, hasta)
        ],
        'articulos_mas_alquilados': [
            {
                'articulo_id': row['articulo_id'],
                'referencia': row['articulo__referencia'],
                'articulo': row['articulo__nombre'],
                'cantidad': _monto(row['cantidad']),
            }
            for row in articulos_mas_alquilados(negocio, desde, hasta)
        ],
        'cartera_credito': [
            {
                'cliente_id': row['cliente_id'],
                'cliente': row['cliente__nombre'] or 'Sin cliente',
                'saldo': _monto(row['saldo']),
            }
            for row in cartera_credito_por_cliente(negocio)
        ],
        'estado_caja_por_turno': [
            {
                **row,
                'efectivo_esperado': _monto(row['efectivo_esperado']),
                'efectivo_contado': _monto(row['efectivo_contado']),
                'diferencia': _monto(row['diferencia']),
            }
            for row in estado_caja_por_turno(negocio, desde, hasta)
        ],
    }


class CategoriaGastoListView(generics.ListAPIView):
    permission_classes = [IsAuthenticated, IsAdminOrCajero]
    serializer_class = CategoriaGastoSerializer

    def get_queryset(self):
        return CategoriaGasto.objects.filter(
            negocio=self.request.user.negocio, activo=True,
        ).order_by('nombre')


class GastoListCreateView(generics.ListCreateAPIView):
    permission_classes = [IsAuthenticated, IsAdminOrCajero]
    serializer_class = GastoSerializer

    def get_queryset(self):
        queryset = Gasto.objects.filter(
            negocio=self.request.user.negocio,
        ).select_related('categoria', 'turno').order_by('-fecha', '-pk')
        desde = self.request.query_params.get('desde')
        hasta = self.request.query_params.get('hasta')
        for name, value in (('desde', desde), ('hasta', hasta)):
            if value and parse_date(value) is None:
                raise ValidationError({name: 'Usa una fecha válida en formato AAAA-MM-DD.'})
        if desde:
            queryset = queryset.filter(fecha__date__gte=parse_date(desde))
        if hasta:
            queryset = queryset.filter(fecha__date__lte=parse_date(hasta))
        categoria = self.request.query_params.get('categoria')
        if categoria:
            if not categoria.isdigit():
                raise ValidationError({'categoria': 'Debe ser un identificador válido.'})
            queryset = queryset.filter(categoria_id=categoria)
        return queryset

    def create(self, request, *args, **kwargs):
        serializer = RegistrarGastoSerializer(
            data=request.data, context={'request': request},
        )
        serializer.is_valid(raise_exception=True)
        try:
            gasto = registrar_gasto(
                negocio=request.user.negocio,
                usuario=request.user,
                **serializer.validated_data,
            )
        except ValueError as exc:
            raise ValidationError({'detail': str(exc)}) from exc
        return Response(
            GastoSerializer(gasto, context={'request': request}).data,
            status=status.HTTP_201_CREATED,
        )


class ReporteView(APIView):
    permission_classes = [IsAuthenticated, IsAdmin]

    def get(self, request):
        _, _, desde, hasta = _rango_fechas(request, requerido=True)
        return Response(_reporte(request.user.negocio, desde, hasta))


class ReporteExcelView(APIView):
    permission_classes = [IsAuthenticated, IsAdmin]

    def get(self, request):
        desde_dia, hasta_dia, desde, hasta = _rango_fechas(request, requerido=True)
        reporte = _reporte(request.user.negocio, desde, hasta)
        workbook = Workbook()
        sheets = [
            ('Resumen', ['Indicador', 'Valor'], [
                [key.replace('_', ' ').title(), Decimal(value)]
                for key, value in reporte['resumen'].items()
            ]),
            ('Gastos por categoría', ['Categoría', 'Total'], [
                [row['categoria'], Decimal(row['total'])]
                for row in reporte['gastos_por_categoria']
            ]),
            ('Medios de pago', ['Medio de pago', 'Ingresos', 'Egresos'], [
                [row['medio_pago'], Decimal(row['ingresos']), Decimal(row['egresos'])]
                for row in reporte['ingresos_egresos_por_medio']
            ]),
            ('Productos vendidos', ['Referencia', 'Producto', 'Cantidad', 'Ingresos'], [
                [row['referencia'], row['producto'], Decimal(row['cantidad']), Decimal(row['ingresos'])]
                for row in reporte['productos_mas_vendidos']
            ]),
            ('Artículos alquilados', ['Referencia', 'Artículo', 'Cantidad'], [
                [row['referencia'], row['articulo'], Decimal(row['cantidad'])]
                for row in reporte['articulos_mas_alquilados']
            ]),
            ('Cartera de crédito', ['Cliente', 'Saldo'], [
                [row['cliente'], Decimal(row['saldo'])] for row in reporte['cartera_credito']
            ]),
            ('Estado de caja', [
                'Turno', 'Caja', 'Usuario', 'Apertura', 'Estado',
                'Efectivo esperado', 'Efectivo contado', 'Diferencia',
            ], [
                [
                    row['turno_id'], row['caja'], row['usuario'],
                    row['apertura_en'].isoformat(),
                    row['estado'],
                    Decimal(row['efectivo_esperado']),
                    Decimal(row['efectivo_contado']) if row['efectivo_contado'] is not None else None,
                    Decimal(row['diferencia']) if row['diferencia'] is not None else None,
                ]
                for row in reporte['estado_caja_por_turno']
            ]),
        ]
        for index, (name, headers, rows) in enumerate(sheets):
            sheet = workbook.active if index == 0 else workbook.create_sheet()
            sheet.title = name
            sheet.append(headers)
            for row in rows:
                sheet.append(row)
            sheet.freeze_panes = 'A2'
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='254E70')
            for column in sheet.columns:
                width = min(max(len(str(cell.value or '')) for cell in column) + 2, 42)
                sheet.column_dimensions[column[0].column_letter].width = width

        output = BytesIO()
        workbook.save(output)
        filename = f'reporte-financiero-{desde_dia:%Y%m%d}-{hasta_dia:%Y%m%d}.xlsx'
        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response


def _alertas(negocio, usuario):
    alertas = []
    for producto in productos_bajo_minimo(negocio)[:100]:
        alertas.append({
            'clave': f'stock-{producto.pk}',
            'titulo': 'Producto agotado' if producto.stock_actual == 0 else 'Stock bajo',
            'detalle': (
                f'{producto.referencia} · {producto.nombre} · '
                f'quedan {producto.stock_actual}'
            ),
            'tipo': 'danger' if producto.stock_actual == 0 else 'warning',
        })
    for alquiler in alquileres_vencidos(negocio)[:100]:
        alertas.append({
            'clave': f'rental-{alquiler.pk}',
            'titulo': f'Alquiler #{alquiler.pk} vencido',
            'detalle': f'{alquiler.cliente.nombre} · debía devolverse '
            f'{timezone.localtime(alquiler.fecha_prevista_devolucion):%d/%m/%Y %H:%M}',
            'tipo': 'danger',
        })
    creditos = cartera_credito_por_cliente(negocio)
    for row in creditos[:100]:
        alertas.append({
            'clave': f'credit-{row["cliente_id"]}',
            'titulo': 'Cliente con saldo pendiente',
            'detalle': f'{row["cliente__nombre"] or "Sin cliente"} · saldo ${row["saldo"]}',
            'tipo': 'warning',
        })
    separados = Venta.objects.filter(
        negocio=negocio, tipo=Venta.Tipo.SEPARADO,
        estado=Venta.Estado.PENDIENTE, saldo_pendiente__gt=0,
    ).select_related('cliente').order_by('fecha')[:100]
    for venta in separados:
        alertas.append({
            'clave': f'separate-{venta.pk}',
            'titulo': f'Separado #{venta.numero} con saldo',
            'detalle': f'{venta.cliente.nombre if venta.cliente_id else "Sin cliente"} · '
            f'debe ${venta.saldo_pendiente}',
            'tipo': 'warning',
        })
    turnos = TurnoCaja.objects.filter(
        negocio=negocio, estado=TurnoCaja.Estado.CERRADO,
        diferencia__isnull=False,
    ).exclude(diferencia=0).select_related('caja', 'usuario').order_by('-cierre_en')[:100]
    for turno in turnos:
        alertas.append({
            'clave': f'cash-{turno.pk}',
            'titulo': f'Diferencia en cierre · {turno.caja.nombre}',
            'detalle': f'{turno.usuario.get_full_name() or turno.usuario.username} · '
            f'${turno.diferencia}',
            'tipo': 'danger',
        })

    leidas = set(AlertaLeida.objects.filter(
        negocio=negocio, usuario=usuario,
        clave__in=[alerta['clave'] for alerta in alertas],
    ).values_list('clave', flat=True))
    ahora = timezone.localtime()
    return [
        {
            'id': alerta['clave'],
            'titulo': alerta['titulo'],
            'detalle': alerta['detalle'],
            'tipo': alerta['tipo'],
            'creada_en': ahora.isoformat(),
            'leida': alerta['clave'] in leidas,
        }
        for alerta in alertas
    ]


class NotificacionesView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrCajero]

    def get(self, request):
        return Response({'results': _alertas(request.user.negocio, request.user)})


class NotificacionesLeidasView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrCajero]

    def post(self, request):
        claves = request.data.get('claves')
        if not isinstance(claves, list) or len(claves) > 200 or any(
            not isinstance(clave, str) or not clave or len(clave) > 160
            for clave in claves
        ):
            raise ValidationError({'claves': 'Envía una lista válida de claves de alertas.'})
        activas = {
            alerta['id']
            for alerta in _alertas(request.user.negocio, request.user)
        }
        if not set(claves).issubset(activas):
            raise ValidationError({'claves': 'Una o más alertas ya no están activas.'})
        AlertaLeida.objects.bulk_create(
            [
                AlertaLeida(
                    negocio=request.user.negocio,
                    creado_por=request.user,
                    usuario=request.user,
                    clave=clave,
                )
                for clave in set(claves)
            ],
            ignore_conflicts=True,
        )
        return Response(status=status.HTTP_204_NO_CONTENT)


class DashboardResumenView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrCajero]

    def get(self, request):
        negocio = request.user.negocio
        hoy = timezone.localdate()
        inicio_hoy = timezone.make_aware(datetime.combine(hoy, time.min))
        inicio_manana = inicio_hoy + timedelta(days=1)
        ventas_hoy = Venta.objects.filter(
            negocio=negocio, fecha__gte=inicio_hoy, fecha__lt=inicio_manana,
        ).exclude(
            estado__in=[Venta.Estado.ANULADA, Venta.Estado.CANCELADA],
        ).aggregate(total=Sum('total'))['total'] or Decimal('0')

        turno = TurnoCaja.objects.filter(
            negocio=negocio, usuario=request.user, estado=TurnoCaja.Estado.ABIERTO,
        ).select_related('caja').first()
        resumen_turno_actual = resumen_turno(turno) if turno else None
        inicio_semana = hoy - timedelta(days=6)
        inicio_grafica = timezone.make_aware(datetime.combine(inicio_semana, time.min))
        ventas_por_dia = {
            row['dia']: row['total'] or Decimal('0')
            for row in Venta.objects.filter(
                negocio=negocio, fecha__gte=inicio_grafica, fecha__lt=inicio_manana,
            ).exclude(
                estado__in=[Venta.Estado.ANULADA, Venta.Estado.CANCELADA],
            ).annotate(dia=TruncDate('fecha')).values('dia').annotate(total=Sum('total'))
        }
        gastos_por_dia = {
            row['dia']: row['total'] or Decimal('0')
            for row in Gasto.objects.filter(
                negocio=negocio, fecha__gte=inicio_grafica, fecha__lt=inicio_manana,
            ).annotate(dia=TruncDate('fecha')).values('dia').annotate(total=Sum('valor'))
        }
        alertas = _alertas(negocio, request.user)
        response = {
            'ventas_hoy': str(ventas_hoy),
            'caja_abierta': turno is not None,
            'saldo_caja': str(resumen_turno_actual['efectivo_esperado'])
            if resumen_turno_actual else '0',
            'turno': turno.caja.nombre if turno else None,
            'alertas': alertas[:8],
            'conteo_alertas': sum(not alerta['leida'] for alerta in alertas),
            'grafica': [
                {
                    'dia': (inicio_semana + timedelta(days=offset)).isoformat(),
                    'ventas': str(ventas_por_dia.get(inicio_semana + timedelta(days=offset), 0)),
                    'gastos': str(gastos_por_dia.get(inicio_semana + timedelta(days=offset), 0)),
                }
                for offset in range(7)
            ],
        }
        if request.user.is_superuser or request.user.rol == 'ADMIN':
            inicio_mes = timezone.make_aware(
                datetime.combine(hoy.replace(day=1), time.min),
            )
            fin_mes = inicio_manana
            resultado = resumen_periodo(negocio, inicio_mes, fin_mes)
            response['inversion'] = _monto(ProductoVenta.objects.filter(
                negocio=negocio, activo=True,
            ).aggregate(total=Sum(ExpressionWrapper(
                F('stock_actual') * F('costo_promedio'),
                output_field=ProductoVenta._meta.get_field('costo_promedio'),
            )))['total'] or Decimal('0'))
            response['ganancia_neta'] = _monto(resultado['ganancia_neta'])
            response['gastos_periodo'] = _monto(resultado['gastos_operativos'])
        return Response(response)
