from datetime import timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from django.contrib.auth import get_user_model
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from alquileres.models import Alquiler, ArticuloAlquiler, DetalleAlquiler
from cajas.models import Caja, MovimientoCaja
from cajas.services import abrir_turno, registrar_movimiento
from core.models import Cliente, Negocio
from finanzas.models import CategoriaGasto, CompraInventario, Gasto
from inventario.models import ProductoVenta
from ventas.models import DetalleVenta, Venta

User = get_user_model()


@pytest.fixture
def escenario_reportes(db):
    negocio = Negocio.objects.create(nombre='Negocio financiero')
    admin = User.objects.create_user(
        username='admin_finanzas', password='Segura123456!',
        negocio=negocio, rol='ADMIN',
    )
    cajero = User.objects.create_user(
        username='cajero_finanzas', password='Segura123456!',
        negocio=negocio, rol='CAJERO',
    )
    categoria = CategoriaGasto.objects.create(negocio=negocio, nombre='Operación')
    producto = ProductoVenta.objects.create(
        negocio=negocio, referencia='PROD-01', nombre='Producto prueba',
        costo_promedio='100.00', precio_venta='250.00', stock_actual='8',
        stock_minimo='2',
    )
    compra = CompraInventario.objects.create(
        negocio=negocio, creado_por=admin, numero=1, total='300.00',
    )
    venta = Venta.objects.create(
        negocio=negocio, creado_por=cajero, numero=1, tipo=Venta.Tipo.CONTADO,
        estado=Venta.Estado.PAGADA, subtotal='500.00', total='500.00',
        saldo_pendiente='0',
    )
    DetalleVenta.objects.create(
        negocio=negocio, creado_por=cajero, venta=venta, producto=producto,
        cantidad='2', precio_unitario='250.00', costo_unitario='100.00',
    )
    cliente = Cliente.objects.create(negocio=negocio, nombre='Cliente cartera')
    Venta.objects.create(
        negocio=negocio, creado_por=cajero, numero=2, cliente=cliente,
        tipo=Venta.Tipo.CREDITO, estado=Venta.Estado.PENDIENTE,
        subtotal='150.00', total='150.00', saldo_pendiente='150.00',
    )
    articulo = ArticuloAlquiler.objects.create(
        negocio=negocio, referencia='ALQ-01', nombre='Andamio',
        tipo=ArticuloAlquiler.Tipo.ANDAMIO, cantidad_total='4',
        cantidad_disponible='3', tarifa_diaria='100.00', costo_diario='10.00',
    )
    salida = timezone.now()
    alquiler = Alquiler.objects.create(
        negocio=negocio, creado_por=cajero, cliente=cliente,
        fecha_salida=salida, fecha_prevista_devolucion=salida + timedelta(days=2),
        estado=Alquiler.Estado.ACTIVO, total='200.00',
    )
    DetalleAlquiler.objects.create(
        negocio=negocio, creado_por=cajero, alquiler=alquiler, articulo=articulo,
        cantidad='1', tarifa_dia='100.00',
    )
    gasto = Gasto.objects.create(
        negocio=negocio, creado_por=cajero, categoria=categoria, valor='80.00',
        descripcion='Gasto de prueba', fecha=timezone.now(), medio_pago='EFECTIVO',
    )
    caja = Caja.objects.create(negocio=negocio, creado_por=admin, nombre='Principal')
    turno = abrir_turno(
        negocio=negocio, usuario=cajero, caja=caja, base_inicial='100.00',
    )
    registrar_movimiento(
        negocio=negocio, usuario=cajero, turno=turno,
        tipo=MovimientoCaja.Tipo.INGRESO, concepto='Ingreso de prueba',
        medio_pago=MovimientoCaja.MedioPago.EFECTIVO, valor='450.00',
    )
    return locals()


@pytest.mark.django_db
def test_reporte_calcula_resultado_a_mano_y_excel_valido(escenario_reportes):
    data = escenario_reportes
    client = APIClient()
    client.force_authenticate(data['admin'])
    dia = timezone.localdate().isoformat()
    response = client.get(f'/api/reportes/?desde={dia}&hasta={dia}')

    assert response.status_code == 200, response.data
    # Ventas 500 + crédito 150 + alquileres 200 - costos 200 y 20 - gastos 80.
    assert response.data['resumen'] == {
        'inversion_compras': '300.00',
        'ingresos_ventas': '650.00',
        'costo_ventas': '200.00',
        'ingresos_alquileres': '200.00',
        'costo_alquiler_estimado': '20.00',
        'ganancia_bruta': '630.00',
        'gastos_operativos': '80.00',
        'ganancia_neta': '550.00',
    }
    assert response.data['gastos_por_categoria'][0]['total'] == '80.00'
    assert response.data['ingresos_egresos_por_medio'][0]['ingresos'] == '450.00'
    assert response.data['productos_mas_vendidos'][0]['cantidad'] == '2.00'
    assert response.data['articulos_mas_alquilados'][0]['cantidad'] == '1.00'
    assert response.data['cartera_credito'][0]['saldo'] == '150.00'

    excel = client.get(f'/api/reportes/exportar-excel/?desde={dia}&hasta={dia}')
    assert excel.status_code == 200
    workbook = load_workbook(BytesIO(excel.content), read_only=True, data_only=True)
    assert workbook.sheetnames == [
        'Resumen', 'Gastos por categoría', 'Medios de pago', 'Productos vendidos',
        'Artículos alquilados', 'Cartera de crédito', 'Estado de caja',
    ]
    rows = list(workbook['Resumen'].iter_rows(values_only=True))
    assert ('Ganancia Neta', Decimal('550.00')) in rows


@pytest.mark.django_db
def test_reportes_solo_admin_y_dashboard_oculta_inversion_y_ganancia(escenario_reportes):
    data = escenario_reportes
    client = APIClient()
    dia = timezone.localdate().isoformat()
    client.force_authenticate(data['cajero'])
    assert client.get(f'/api/reportes/?desde={dia}&hasta={dia}').status_code == 403
    assert client.get(f'/api/reportes/exportar-excel/?desde={dia}&hasta={dia}').status_code == 403
    summary = client.get('/api/dashboard/resumen/')
    assert summary.status_code == 200
    assert 'inversion' not in summary.data
    assert 'ganancia_neta' not in summary.data

    client.force_authenticate(data['admin'])
    assert client.get(f'/api/reportes/?desde={dia}&hasta={dia}').status_code == 200
    admin_summary = client.get('/api/dashboard/resumen/')
    assert admin_summary.status_code == 200, admin_summary.data
    assert admin_summary.data['inversion'] == '800.00'
    assert admin_summary.data['ganancia_neta'] == '550.00'


@pytest.mark.django_db
def test_api_gastos_fuera_de_caja_y_notificacion_leida(escenario_reportes):
    data = escenario_reportes
    client = APIClient()
    client.force_authenticate(data['cajero'])
    category = client.get('/api/gastos/categorias/')
    assert category.status_code == 200
    gasto = client.post('/api/gastos/', {
        'categoria': data['categoria'].pk,
        'valor': '25.00',
        'descripcion': 'Servicio externo',
        'fecha': timezone.now().isoformat(),
        'medio_pago': 'TRANSFERENCIA',
        'desde_caja': False,
    }, format='json')
    assert gasto.status_code == 201, gasto.data
    assert gasto.data['turno'] is None
    assert not MovimientoCaja.objects.filter(gasto_id=gasto.data['id']).exists()

    alerts = client.get('/api/notificaciones/alertas/')
    assert alerts.status_code == 200
    assert alerts.data['results']
    key = alerts.data['results'][0]['id']
    marked = client.post('/api/notificaciones/marcar-leidas/', {'claves': [key]}, format='json')
    assert marked.status_code == 204
    refreshed = client.get('/api/notificaciones/alertas/')
    assert next(item for item in refreshed.data['results'] if item['id'] == key)['leida']
