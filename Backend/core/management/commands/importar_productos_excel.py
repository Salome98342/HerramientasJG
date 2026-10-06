from decimal import Decimal
from django.core.management.base import BaseCommand, CommandError
from openpyxl import load_workbook
from core.models import Negocio, Usuario
from inventario.models import CategoriaProducto, ProductoVenta, MovimientoInventario

class Command(BaseCommand):
    help='Importa productos desde Excel. Plantilla provisional; ajustar columnas cuando el cliente entregue su archivo.'
    def add_arguments(self, parser):
        parser.add_argument('archivo')
        parser.add_argument('--negocio-id', type=int, required=True)
        parser.add_argument('--usuario-id', type=int, required=True)
        parser.add_argument('--dry-run', action='store_true')
    def handle(self,*args,**opts):
        negocio=Negocio.objects.get(pk=opts['negocio_id']); usuario=Usuario.objects.get(pk=opts['usuario_id'])
        wb=load_workbook(opts['archivo'], read_only=True, data_only=True); ws=wb.active
        rows=ws.iter_rows(values_only=True); headers=[str(x or '').strip().lower() for x in next(rows)]
        required={'referencia','nombre'}
        missing=required-set(headers)
        if missing: raise CommandError(f'Faltan columnas obligatorias: {sorted(missing)}')
        idx={h:i for i,h in enumerate(headers)}; creados=actualizados=movimientos=0
        for row in rows:
            referencia=str(row[idx['referencia']] or '').strip(); nombre=str(row[idx['nombre']] or '').strip()
            if not referencia: continue
            categoria=None
            if 'categoria' in idx and row[idx['categoria']]:
                categoria,_=CategoriaProducto.objects.get_or_create(negocio=negocio,nombre=str(row[idx['categoria']]).strip(),defaults={'creado_por':usuario})
            defaults={'nombre':nombre,'categoria':categoria,'creado_por':usuario}
            for field in ['costo_promedio','precio_venta','stock_minimo']:
                if field in idx and row[idx[field]] is not None: defaults[field]=Decimal(str(row[idx[field]]))
            producto=ProductoVenta.objects.filter(negocio=negocio,referencia=referencia).first()
            stock_nuevo=Decimal(str(row[idx['stock_actual']])) if 'stock_actual' in idx and row[idx['stock_actual']] is not None else None
            if opts['dry_run']: continue
            if not producto:
                producto=ProductoVenta.objects.create(negocio=negocio,referencia=referencia,stock_actual=Decimal('0'),**defaults); creados+=1
            else:
                for k,v in defaults.items(): setattr(producto,k,v)
                producto.save(); actualizados+=1
            if stock_nuevo is not None and stock_nuevo != producto.stock_actual:
                diferencia=stock_nuevo-producto.stock_actual
                MovimientoInventario.objects.create(negocio=negocio,creado_por=usuario,producto=producto,tipo=MovimientoInventario.Tipo.AJUSTE,
                    direccion=MovimientoInventario.Direccion.ENTRADA if diferencia>0 else MovimientoInventario.Direccion.SALIDA,cantidad=abs(diferencia),
                    costo_unitario=producto.costo_promedio,motivo='Importación Excel',usuario=usuario,origen_tipo='IMPORTACION_EXCEL')
                producto.stock_actual=stock_nuevo; producto.save(update_fields=['stock_actual','actualizado_en']); movimientos+=1
        self.stdout.write(self.style.SUCCESS(f'Importación: {creados} creados, {actualizados} actualizados, {movimientos} ajustes de stock.'))
