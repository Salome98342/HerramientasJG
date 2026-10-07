from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from io import BytesIO
from zipfile import BadZipFile

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from alquileres.models import ArticuloAlquiler, MovimientoInventarioAlquiler
from core.models import Negocio, Usuario
from inventario.models import CategoriaProducto, MovimientoInventario, ProductoVenta

MAX_IMPORT_BYTES = 10 * 1024 * 1024

SALE_HEADERS = (
    'referencia', 'nombre', 'categoria', 'costo_promedio', 'precio_venta',
    'stock_actual', 'stock_minimo', 'activo',
)
RENTAL_HEADERS = (
    'referencia', 'nombre', 'tipo', 'cantidad_total', 'tarifa_diaria',
    'costo_diario', 'valor_reposicion', 'activo',
)
RENTAL_TYPES = {value for value, _label in ArticuloAlquiler.Tipo.choices}
TRUE_VALUES = {'si', 'sí', 'true', '1', 'activo'}
FALSE_VALUES = {'no', 'false', '0', 'inactivo'}


@dataclass
class ImportRow:
    sheet: str
    number: int
    reference: str
    values: dict
    errors: list[str]
    existing: bool = False

    def as_dict(self):
        if self.errors:
            status = 'ERROR'
        elif self.existing:
            status = 'OMITIR_EXISTENTE'
        else:
            status = 'CREAR'
        return {
            'hoja': self.sheet,
            'fila': self.number,
            'referencia': self.reference,
            'estado': status,
            'errores': self.errors,
        }


def crear_plantilla_excel():
    workbook = Workbook()
    sales = workbook.active
    sales.title = 'Venta'
    sales.append(SALE_HEADERS)
    sales.freeze_panes = 'A2'
    rentals = workbook.create_sheet('Alquiler')
    rentals.append(RENTAL_HEADERS)
    rentals.freeze_panes = 'A2'
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _text(value):
    return str(value).strip() if value is not None else ''


def _decimal(value, field_name, errors, *, required=False):
    if value in (None, ''):
        if required:
            errors.append(f'La columna {field_name} es obligatoria.')
        return Decimal('0')
    try:
        result = Decimal(str(value).strip())
        if not result.is_finite() or result < 0:
            raise InvalidOperation
        if result.as_tuple().exponent < -2:
            errors.append(f'{field_name} admite máximo dos decimales.')
        if result > Decimal('999999999999.99'):
            errors.append(f'{field_name} supera el máximo permitido (999999999999.99).')
        return result
    except (InvalidOperation, ValueError, TypeError):
        errors.append(f'{field_name} debe ser un número mayor o igual a cero.')
        return Decimal('0')


def _boolean(value, field_name, errors):
    if value in (None, ''):
        return True
    if isinstance(value, bool):
        return value
    normalized = _text(value).lower()
    if normalized in TRUE_VALUES:
        return True
    if normalized in FALSE_VALUES:
        return False
    errors.append(f'{field_name} debe ser Sí o No.')
    return True


def _read_sheet(worksheet, expected_headers, business, seen):
    rows = worksheet.iter_rows(values_only=True)
    header_row = next(rows, None)
    if not header_row:
        return [], [f'La hoja {worksheet.title} está vacía; debe incluir encabezados en la fila 1.']
    headers = [_text(value).lower() for value in header_row]
    indexes = {header: index for index, header in enumerate(headers) if header}
    missing = [header for header in expected_headers if header not in indexes]
    if missing:
        return [], [
            f'La hoja {worksheet.title} no incluye las columnas requeridas: {", ".join(missing)}.'
        ]

    parsed = []
    structural_errors = []
    existing_references = set()
    model = ProductoVenta if worksheet.title == 'Venta' else ArticuloAlquiler
    existing_references.update(
        model.objects.filter(negocio=business).values_list('referencia', flat=True)
    )
    for number, row in enumerate(rows, start=2):
        cells = {header: row[index] if index < len(row) else None for header, index in indexes.items()}
        if all(value in (None, '') for value in cells.values()):
            continue

        errors = []
        reference = _text(cells.get('referencia'))
        name = _text(cells.get('nombre'))
        if not reference:
            errors.append('La referencia es obligatoria.')
        elif len(reference) > 60:
            errors.append('La referencia admite máximo 60 caracteres.')
        if not name:
            errors.append('El nombre es obligatorio.')
        elif len(name) > 180:
            errors.append('El nombre admite máximo 180 caracteres.')

        if reference in seen[worksheet.title]:
            errors.append('La referencia está repetida en el archivo.')
        elif reference:
            seen[worksheet.title].add(reference)

        if worksheet.title == 'Venta':
            category = _text(cells.get('categoria'))
            if len(category) > 100:
                errors.append('La categoría admite máximo 100 caracteres.')
            values = {
                'referencia': reference,
                'nombre': name,
                'categoria': category,
                'costo_promedio': _decimal(cells.get('costo_promedio'), 'costo_promedio', errors),
                'precio_venta': _decimal(cells.get('precio_venta'), 'precio_venta', errors, required=True),
                'stock_actual': _decimal(cells.get('stock_actual'), 'stock_actual', errors),
                'stock_minimo': _decimal(cells.get('stock_minimo'), 'stock_minimo', errors),
                'activo': _boolean(cells.get('activo'), 'activo', errors),
            }
        else:
            item_type = _text(cells.get('tipo')).upper()
            if item_type not in RENTAL_TYPES:
                errors.append(f'Tipo inválido. Usa: {", ".join(sorted(RENTAL_TYPES))}.')
            values = {
                'referencia': reference,
                'nombre': name,
                'tipo': item_type,
                'cantidad_total': _decimal(cells.get('cantidad_total'), 'cantidad_total', errors, required=True),
                'tarifa_diaria': _decimal(cells.get('tarifa_diaria'), 'tarifa_diaria', errors, required=True),
                'costo_diario': _decimal(cells.get('costo_diario'), 'costo_diario', errors),
                'valor_reposicion': _decimal(cells.get('valor_reposicion'), 'valor_reposicion', errors),
                'activo': _boolean(cells.get('activo'), 'activo', errors),
            }

        parsed.append(ImportRow(
            sheet=worksheet.title,
            number=number,
            reference=reference,
            values=values,
            errors=errors,
            existing=reference in existing_references if reference else False,
        ))
    return parsed, structural_errors


def validar_archivo_excel(content, business):
    if not content:
        raise ValueError('Selecciona un archivo Excel para continuar.')
    if len(content) > MAX_IMPORT_BYTES:
        raise ValueError('El archivo supera el tamaño máximo permitido de 10 MB.')
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError, TypeError) as exc:
        raise ValueError('No se pudo leer el archivo. Usa una plantilla Excel .xlsx válida.') from exc

    try:
        missing_sheets = [name for name in ('Venta', 'Alquiler') if name not in workbook.sheetnames]
        if missing_sheets:
            raise ValueError(f'Faltan las hojas requeridas: {", ".join(missing_sheets)}.')
        rows = []
        errors = []
        seen = {'Venta': set(), 'Alquiler': set()}
        for name, headers in (('Venta', SALE_HEADERS), ('Alquiler', RENTAL_HEADERS)):
            sheet_rows, sheet_errors = _read_sheet(workbook[name], headers, business, seen)
            rows.extend(sheet_rows)
            errors.extend(sheet_errors)
        if not rows:
            errors.append('El archivo no contiene filas de productos o artículos para importar.')
        errors.extend(
            f"Hoja {row.sheet}, fila {row.number}: {message}"
            for row in rows
            for message in row.errors
        )
        return rows, errors
    finally:
        workbook.close()


@transaction.atomic
def confirmar_importacion(*, business: Negocio, user: Usuario, rows: list[ImportRow]):
    imported = {'Venta': 0, 'Alquiler': 0}
    skipped = {'Venta': 0, 'Alquiler': 0}
    movements = {'Venta': 0, 'Alquiler': 0}
    for row in rows:
        if row.errors:
            raise ValueError('No se puede confirmar una importación con filas inválidas.')
        if row.sheet == 'Venta':
            values = row.values.copy()
            category_name = values.pop('categoria')
            category = None
            if category_name:
                category, _ = CategoriaProducto.objects.get_or_create(
                    negocio=business,
                    nombre=category_name,
                    defaults={'creado_por': user},
                )
            product, created = ProductoVenta.objects.get_or_create(
                negocio=business,
                referencia=values['referencia'],
                defaults={
                    **values,
                    'categoria': category,
                    'creado_por': user,
                },
            )
            if not created:
                skipped['Venta'] += 1
                continue
            imported['Venta'] += 1
            if product.stock_actual > 0:
                MovimientoInventario.objects.create(
                    negocio=business,
                    creado_por=user,
                    usuario=user,
                    producto=product,
                    tipo=MovimientoInventario.Tipo.AJUSTE,
                    direccion=MovimientoInventario.Direccion.ENTRADA,
                    cantidad=product.stock_actual,
                    costo_unitario=product.costo_promedio,
                    motivo='Existencia inicial por importación Excel',
                    origen_tipo='IMPORTACION_EXCEL',
                )
                movements['Venta'] += 1
        else:
            values = row.values.copy()
            quantity = values.pop('cantidad_total')
            article, created = ArticuloAlquiler.objects.get_or_create(
                negocio=business,
                referencia=values['referencia'],
                defaults={
                    **values,
                    'cantidad_total': quantity,
                    'cantidad_disponible': quantity,
                    'creado_por': user,
                },
            )
            if not created:
                skipped['Alquiler'] += 1
                continue
            imported['Alquiler'] += 1
            if article.cantidad_total > 0:
                MovimientoInventarioAlquiler.objects.create(
                    negocio=business,
                    creado_por=user,
                    articulo=article,
                    tipo=MovimientoInventarioAlquiler.Tipo.INICIAL,
                    cantidad=article.cantidad_total,
                    motivo='Existencia inicial por importación Excel',
                )
                movements['Alquiler'] += 1
    return {'importados': imported, 'omitidos': skipped, 'movimientos': movements}
